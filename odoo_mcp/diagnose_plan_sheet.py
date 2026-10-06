"""Diagnostic: see the raw structure of Payment Plan Structure sheet."""

from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"C:\Parkgroup Data")
PLAN_WB = ROOT / "Consolidated_Sales_Workbook (4).xlsx"

def find_header(rows, scan: int = 12) -> int:
    best_i, best_n = 0, -1
    for i, r in enumerate(rows[:scan]):
        n = sum(1 for c in r if c not in (None, ""))
        if n > best_n:
            best_n, best_i = n, i
    return best_i

wb = load_workbook(PLAN_WB, read_only=True, data_only=True)
ws = wb["Payment Plan Structure"]
rows = [r for r in ws.iter_rows(values_only=True)]
wb.close()
hi = find_header(rows)
print(f"Header row: {hi + 1}")
hdr = [str(c).strip() if c not in (None, "") else "" for c in rows[hi]]
print(f"Headers: {hdr}")
print()
print("First 30 rows after header:")
for i, r in enumerate(rows[hi + 1:hi + 1 + 30]):
    if not r:
        continue
    # Show first 3 cells
    cells = []
    for j in range(min(3, len(r))):
        val = r[j]
        cells.append(f"{hdr[j] if j < len(hdr) else f'col{j}'}={repr(val)}")
    print(f"{hi + 2 + i:3}: {' | '.join(cells)}")
    # If the first cell looks like a new major section, note it
    if len(r) > 0 and r[0]:
        first = str(r[0]).strip()
        if first.startswith("Project=") or ("▼" in first and "(" in first):
            print(f"     ^^^ section header")