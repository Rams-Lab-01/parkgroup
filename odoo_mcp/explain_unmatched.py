"""Explain which sheet rows fail to find their Odoo unit.

The reconciliation reports 81 sheet rows needing attention, concentrated in
PRY (56). This prints the actual (raw, normalised) unit identifiers so the
cause is visible rather than guessed at.
"""

from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from pgre_client import OdooClient
from reconcile import find_header, norm_unit

for _s in ("stdout", "stderr"):
    _st = getattr(sys, _s, None)
    if _st and hasattr(_st, "reconfigure"):
        try:
            _st.reconfigure(encoding="utf-8", errors="replace")
        except (OSError, ValueError):
            pass

SHEET = Path(r"C:\Parkgroup Data\05-10-2026 Final Full Report.xlsx")
PMAP = {"PBR1": "PARK BEACH RESIDENCE", "PBR2": "PARK BEACH RESIDENCE II",
        "PGV": "Park Golf View Residence", "PRY": "Park Residency"}


def main() -> int:
    c = OdooClient()
    props = c.search_read("product.product", [("is_property", "=", True)],
                          ["name", "default_code", "building_name", "floor",
                           "property_area"], limit=0, order="id asc")
    by_building: dict[str, dict[str, dict]] = defaultdict(dict)
    for p in props:
        b = str(p.get("building_name") or "")
        key = norm_unit(p.get("default_code") or p.get("name"))
        # keep the first; duplicates within a building are a defect to log
        by_building[b].setdefault(key, p)

    print("=" * 96)
    print("ODOO UNITS PER BUILDING (normalised keys)")
    print("=" * 96)
    for b, m in sorted(by_building.items()):
        n_products = sum(1 for p in props if str(p.get("building_name") or "") == b)
        dup = n_products - len(m)
        print(f"  {b:<32} {len(m):>4} distinct keys from {n_products:>4} products"
              + (f"   <-- {dup} DUPLICATE key(s)" if dup else ""))

    wb = load_workbook(SHEET, read_only=True, data_only=True)
    ws = wb["All Units"]
    rows = [r for r in ws.iter_rows(values_only=True)]
    hi = find_header(rows)
    hdr = [str(c).strip() if c not in (None, "") else "" for c in rows[hi]]
    col = {h.lower(): i for i, h in enumerate(hdr) if h}
    ip, iu, icl, ist = col["project name"], col["unit no"], col["client name"], col["status"]

    data = [r for r in rows[hi + 1:]
            if any(x not in (None, "") for x in r)
            and str(r[ip]).strip().upper() not in ("TOTAL", "TOTALS", "")]

    miss: dict[str, list[tuple[str, str, str, str]]] = defaultdict(list)
    for r in data:
        proj = str(r[ip]).strip().upper()
        raw = r[iu]
        unit = norm_unit(raw)
        if proj not in PMAP or not unit:
            continue
        if unit not in by_building[PMAP[proj]]:
            miss[proj].append((unit, str(raw), str(r[icl] or "").strip(),
                               str(r[ist] or "").strip()))

    print()
    print("=" * 96)
    print("SHEET ROWS WITH NO ODOO UNIT")
    print("=" * 96)
    for proj, items in sorted(miss.items()):
        print(f"\n  {proj} -> {PMAP[proj]}: {len(items)} row(s) unmatched")
        for unit, raw, client, status in items[:20]:
            print(f"     normalised={unit!r:<12} raw={raw!r:<18} status={status!r:<12} client={client[:28]!r}")
        if len(items) > 20:
            print(f"     ... and {len(items) - 20} more")

    # Are the unmatched sheet units "Available" (i.e. expected to be absent)?
    print()
    print("=" * 96)
    print("STATUS BREAKDOWN OF UNMATCHED ROWS")
    print("=" * 96)
    allrows = [i for items in miss.values() for i in items]
    by_status: dict[str, int] = defaultdict(int)
    for _, _, _, s in allrows:
        by_status[s] += 1
    for s, n in sorted(by_status.items(), key=lambda kv: -kv[1]):
        print(f"  {n:>5}  {s!r}")
    print()
    print("  If unmatched rows are overwhelmingly 'Available', absence from Odoo is")
    print("  expected rather than a defect - the sheet lists the whole inventory,")
    print("  while Odoo holds only what was loaded.")
    wb.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())