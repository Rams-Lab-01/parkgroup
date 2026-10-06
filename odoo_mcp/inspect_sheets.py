"""Inspect the workbook structure in C:\\Parkgroup Data.

Prints sheet names, dimensions and a sample of headers/rows per sheet so the
reconciliation engine can be written against the real column layout instead of
guesses. Read-only: nothing is written.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

# Excel content is UTF-8 (Arabic, box-drawing); the Windows console defaults to
# cp1252 and would raise UnicodeEncodeError.
for _stream in ("stdout", "stderr"):
    _s = getattr(sys, _stream, None)
    if _s is not None and hasattr(_s, "reconfigure"):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except (OSError, ValueError):
            pass

ROOT = Path(r"C:\Parkgroup Data")
MAX_SAMPLE_ROWS = 4
MAX_SHEETS_PER_FILE = 12


def cell(v: object) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return s if len(s) <= 34 else s[:31] + "..."


def main() -> int:
    files = sorted(
        [p for p in ROOT.glob("*.xlsx") if not p.name.startswith("~$")],
        key=lambda p: p.name,
    )
    if not files:
        print(f"No .xlsx files found in {ROOT}")
        return 1

    for path in files:
        print("=" * 100)
        print(f"FILE: {path.name}  ({path.stat().st_size:,} bytes)")
        print("=" * 100)
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            print(f"  ERROR opening: {exc}\n")
            continue

        for idx, ws in enumerate(wb.worksheets[:MAX_SHEETS_PER_FILE]):
            # ReadOnlyWorksheet has no .dimensions; only max_row/max_column.
            dims = getattr(ws, "dimensions", None) or f"{ws.max_row}x{ws.max_column}"
            print(f"\n  SHEET: {ws.title!r}  dims={dims} "
                  f"max_row={ws.max_row} max_col={ws.max_column}")
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                rows.append(row)
                if i >= MAX_SAMPLE_ROWS + 2:
                    break
            if not rows:
                print("    (empty)")
                continue
            # Heuristic: the first row with the most non-empty cells is the header.
            header_idx, header = 0, rows[0]
            best = sum(1 for c in header if c not in (None, ""))
            for i, r in enumerate(rows[:4]):
                filled = sum(1 for c in r if c not in (None, ""))
                if filled > best:
                    best, header_idx, header = filled, i, r
            hdr = [cell(c) for c in header]
            print(f"    header row {header_idx + 1}:")
            print("      " + " | ".join(f"{h}" for h in hdr[:28]))
            for r in rows[header_idx + 1:header_idx + 1 + MAX_SAMPLE_ROWS]:
                if not any(c not in (None, "") for c in r):
                    continue
                print("      " + " | ".join(cell(c) for c in r[:28]))
        if len(wb.worksheets) > MAX_SHEETS_PER_FILE:
            print(f"\n  ... and {len(wb.worksheets) - MAX_SHEETS_PER_FILE} more sheet(s)")
        wb.close()
        print()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())