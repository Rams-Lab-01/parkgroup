"""Locate installment / payment-plan data on both sides.

Sheets side
    `All Units` carries milestone columns (10% Payment Date/Amount,
    20% Payment Date/Amount, Escrow %) plus a Payment Plan description.
    `Consolidated_Sales_Workbook (4).xlsx` has dedicated sheets
    `Payment Plan Structure` and `EMI vs Milestone Comparison`.

Odoo side
    Installments are not a separate model - a scheduled installment is an
    `account.move.line` carrying `date_maturity`. The plan template is
    `account.payment.term` / `account.payment.term.line`, referenced from the
    invoice through `invoice_payment_term_id`.

This prints the real shapes so the validator is written against evidence.
"""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from pgre_client import OdooClient, OdooError

for _s in ("stdout", "stderr"):
    _st = getattr(sys, _s, None)
    if _st and hasattr(_st, "reconfigure"):
        try:
            _st.reconfigure(encoding="utf-8", errors="replace")
        except (OSError, ValueError):
            pass

ROOT = Path(r"C:\Parkgroup Data")
PLAN_WORKBOOK = ROOT / "Consolidated_Sales_Workbook (4).xlsx"
FULL_REPORT = ROOT / "05-10-2026 Final Full Report.xlsx"
PLAN_SHEETS = ("Payment Plan Structure", "EMI vs Milestone Comparison",
               "Unit Payment Collection")


def find_header(rows, scan: int = 12) -> int:
    best_i, best_n = 0, -1
    for i, r in enumerate(rows[:scan]):
        n = sum(1 for c in r if c not in (None, ""))
        if n > best_n:
            best_n, best_i = n, i
    return best_i


def dump_sheet(path: Path, sheet: str, max_rows: int = 14) -> None:
    if not path.exists():
        print(f"  (workbook missing: {path.name})")
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        print(f"  (sheet {sheet!r} not in {path.name})")
        wb.close()
        return
    ws = wb[sheet]
    rows = [r for r in ws.iter_rows(values_only=True)]
    hi = find_header(rows)
    hdr = [str(c).strip() if c not in (None, "") else "" for c in rows[hi]]
    data = [r for r in rows[hi + 1:] if any(c not in (None, "") for c in r)]
    print(f"\n  --- {path.name} :: {sheet!r} ---")
    print(f"      header row {hi + 1}, {len(data)} data rows, {len([h for h in hdr if h])} columns")
    print(f"      columns: {[h for h in hdr if h]}")
    for r in data[:max_rows]:
        cells = [f"{hdr[i]}={r[i]}" for i in range(len(hdr))
                 if i < len(r) and hdr[i] and r[i] not in (None, "")]
        print(f"      - {' | '.join(cells)[:300]}")
    wb.close()


def main() -> int:
    print("=" * 100)
    print("A. SHEET-SIDE PAYMENT PLAN DATA")
    print("=" * 100)
    for sheet in PLAN_SHEETS:
        dump_sheet(PLAN_WORKBOOK, sheet)
    dump_sheet(FULL_REPORT, "Unit Payment Collection", max_rows=8)

    c = OdooClient()
    print()
    print("=" * 100)
    print("B. ODOO PAYMENT TERMS (the plan templates)")
    print("=" * 100)
    fg = c.fields_get("account.payment.term", attributes=["type", "string"])
    print(f"  account.payment.term fields: {sorted(fg)[:22]}")
    terms = c.search_read("account.payment.term", [],
                          ["name", "company_id", "active", "note",
                           "early_discount", "discount_percentage"],
                          limit=0, order="id asc")
    print(f"  total payment terms: {len(terms)}")
    used = Counter()
    for m in c.search_all_terms() if hasattr(c, "search_all_terms") else []:
        pass
    for t in terms[:30]:
        print(f"    id={t['id']:<5} {str(t.get('name'))[:56]:<58} "
              f"co={str((t.get('company_id') or ['','-'])[1])[:34]}")

    print()
    print("=" * 100)
    print("C. PAYMENT TERM LINES (the installment definition)")
    print("=" * 100)
    lines = c.search_read("account.payment.term.line", [],
                          ["payment_term_id", "value", "value_amount", "delay_type",
                           "nb_days", "days_next_month"],
                          limit=0, order="payment_term_id asc, id asc")
    print(f"  total term lines: {len(lines)}")
    by_term: dict[int, list[dict]] = {}
    for ln in lines:
        tid = (ln.get("payment_term_id") or [None])[0]
        by_term.setdefault(tid, []).append(ln)
    print(f"  terms that actually have lines: {len(by_term)}")
    term_name = {t["id"]: t["name"] for t in terms}
    for tid, lns in list(by_term.items())[:12]:
        desc = ", ".join(
            f"{l.get('value')}:{l.get('value_amount')}%"
            if l.get("value") == "percent"
            else f"{l.get('value')}:{l.get('value_amount')}"
            for l in lns)
        print(f"    {str(term_name.get(tid))[:40]:<42} {len(lns):>3} lines: {desc[:90]}")

    print()
    print("=" * 100)
    print("D. WHICH INVOICES CARRY A PAYMENT TERM?")
    print("=" * 100)
    groups = c.execute_kw("account.move", "read_group", [], {
        "domain": [("move_type", "in", ["out_invoice", "out_refund"]),
                   ("state", "=", "posted")],
        "fields": ["invoice_payment_term_id"], "groupby": ["invoice_payment_term_id"],
        "lazy": False}) or []
    print(f"  distinct payment terms on posted customer invoices: {len(groups)}")
    for g in sorted(groups, key=lambda x: -(x.get("invoice_payment_term_id_count") or 0))[:15]:
        tid = g.get("invoice_payment_term_id") or [None, "(none)"]
        print(f"    {str(tid[1])[:48]:<50} {g.get('invoice_payment_term_id_count'):>6}")

    print()
    print("=" * 100)
    print("E. ACTUAL SCHEDULED INSTALLMENTS (move lines with date_maturity)")
    print("=" * 100)
    fg = c.fields_get("account.move.line", attributes=["type", "string"])
    for f in ("date_maturity", "amount_residual", "amount_currency", "payment_date",
              "parent_state"):
        if f in fg:
            print(f"  account.move.line.{f:<20} {fg[f].get('type'):<11} {fg[f].get('string')}")
    dom = [("move_id.move_type", "in", ["out_invoice"]),
           ("move_id.state", "=", "posted"),
           ("date_maturity", "!=", False),
           ("display_type", "=", "payment_term")]
    try:
        n_inst = c.search_count("account.move.line", dom)
        print(f"  payable-term installment lines on customer invoices: {n_inst:,}")
    except OdooError as exc:
        print(f"  count failed: {str(exc)[:120]}")
        n_inst = 0
    if n_inst:
        rows = c.search_read("account.move.line", dom,
                             ["move_id", "date_maturity", "amount_currency",
                              "amount_residual", "currency_id", "partner_id",
                              "name"],
                             limit=14, order="move_id desc")
        for r in rows:
            print(f"    {str((r.get('move_id') or ['','?'])[1])[:20]:<22}"
                  f"due={str(r.get('date_maturity')):<12}"
                  f"amt={r.get('amount_currency')!s:>13}"
                  f"residual={r.get('amount_residual')!s:>13}"
                  f" {(r.get('currency_id') or ['','?'])[1]}")

    print()
    print("=" * 100)
    print("F. INSTALLMENTS PER INVOICE (how many milestones each invoice has)")
    print("=" * 100)
    if n_inst:
        sample_ids = [r["id"] for r in c.search_read(
            "account.move", [("move_type", "=", "out_invoice"), ("state", "=", "posted")],
            ["id"], limit=40, order="id desc")]
        counts = Counter()
        for mid in sample_ids:
            n = c.search_count("account.move.line", [
                ("move_id", "=", mid), ("date_maturity", "!=", False),
                ("display_type", "=", "payment_term")])
            counts[n] += 1
        print(f"  installment-count distribution over {len(sample_ids)} recent invoices:")
        for k in sorted(counts):
            print(f"    {k} installment line(s): {counts[k]} invoice(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())