"""Extract the vocabulary the reconciler needs from the workbooks.

Prints, per "All Units"-style sheet: the detected header row, the exact column
labels, and the distinct values of the categorical columns that must map to Odoo
(project name, unit type, status, SPA/RF status). That is what the
sheet -> Odoo mapping spec is built from.
"""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

for _s in ("stdout", "stderr"):
    _st = getattr(sys, _s, None)
    if _st and hasattr(_st, "reconfigure"):
        try:
            _st.reconfigure(encoding="utf-8", errors="replace")
        except (OSError, ValueError):
            pass

ROOT = Path(r"C:\Parkgroup Data")
PRIMARY = ROOT / "05-10-2026 Final Full Report.xlsx"
UNIT_SHEETS = ("All Units", "PGV Units", "PBR1 Units", "PBR2 Units", "PRY Units")
CATEGORICAL_HINTS = ("project", "status", "type", "nationality", "broker",
                     "oqood", "spa", "rf ", "floor")


def find_header(rows: list[tuple], scan: int = 12):
    """Header = the densest row within the first `scan` rows."""
    best_i, best_n = 0, -1
    for i, r in enumerate(rows[:scan]):
        n = sum(1 for c in r if c not in (None, ""))
        if n > best_n:
            best_n, best_i = n, i
    return best_i, best_n


def main() -> int:
    wb = load_workbook(PRIMARY, read_only=True, data_only=True)
    print(f"FILE: {PRIMARY.name}")

    for sheet in UNIT_SHEETS:
        if sheet not in wb.sheetnames:
            print(f"\n  (sheet {sheet!r} not present)")
            continue
        ws = wb[sheet]
        rows = [r for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        hi, filled = find_header(rows)
        header = [str(c).strip() if c not in (None, "") else "" for c in rows[hi]]
        data = [r for r in rows[hi + 1:] if any(c not in (None, "") for c in r)]

        print()
        print("=" * 90)
        print(f"SHEET {sheet!r}: header on row {hi + 1} ({filled} cols), {len(data)} data rows")
        print("=" * 90)
        for i, h in enumerate(header):
            if h:
                print(f"   [{i:>2}] {h}")

        idx = {h.lower(): i for i, h in enumerate(header) if h}
        for h, i in idx.items():
            if any(k in h for k in CATEGORICAL_HINTS):
                vals = Counter(
                    str(r[i]).strip() for r in data
                    if i < len(r) and r[i] not in (None, "")
                )
                if 1 < len(vals) <= 25:
                    print(f"   distinct {h!r}:")
                    for v, n in vals.most_common():
                        print(f"        {n:>4}  {v}")

    # Cross-check: do the per-project sheets agree with 'All Units'?
    if "All Units" in wb.sheetnames and "PBR1 Units" in wb.sheetnames:
        print()
        print("=" * 90)
        print("CROSS-CHECK: All Units vs per-project sheets")
        print("=" * 90)
        for sheet in UNIT_SHEETS:
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            rows = [r for r in ws.iter_rows(values_only=True)]
            hi, _ = find_header(rows)
            header = [str(c).strip() if c not in (None, "") else "" for c in rows[hi]]
            data = [r for r in rows[hi + 1:] if any(c not in (None, "") for c in r)]
            pj = next((i for i, h in enumerate(header)
                       if h.lower() == "project name"), None)
            un = next((i for i, h in enumerate(header)
                       if h.lower() == "unit no"), None)
            projects = Counter(str(r[pj]).strip() for r in data
                               if pj is not None and pj < len(r) and r[pj] not in (None, ""))
            units = [str(r[un]).strip() for r in data
                     if un is not None and un < len(r) and r[un] not in (None, "")]
            print(f"  {sheet:<16} rows={len(data):>4}  projects={dict(projects)}")
            print(f"  {'':<16} unit sample={units[:8]}")
    wb.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())