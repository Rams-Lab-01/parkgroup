"""Report the company-attribution sheets produced by reconcile.py."""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

for _s in ("stdout", "stderr"):
    _st = getattr(sys, _s, None)
    if _st and hasattr(_st, "reconfigure"):
        try:
            _st.reconfigure(encoding="utf-8", errors="replace")
        except (OSError, ValueError):
            pass

WB = Path(__file__).parent / "reconciliation_2026-10-06.xlsx"


def rows_of(ws):
    rows = list(ws.iter_rows(values_only=True))
    return (rows[0], rows[1:]) if rows else ([], [])


def main() -> int:
    wb = load_workbook(WB, read_only=True, data_only=True)
    print("SHEETS:", wb.sheetnames)

    if "Company_Map" in wb.sheetnames:
        hdr, body = rows_of(wb["Company_Map"])
        print()
        print("=" * 104)
        print(f"COMPANY_MAP  ({len(body)} rows)")
        print("=" * 104)
        for r in body:
            d = dict(zip(hdr, r))
            print(f"  {str(d.get('odoo_building')):<32} units={str(d.get('units')):<5} "
                  f"owners={str(d.get('distinct_owners')):<3} "
                  f"mixed={str(d.get('mixed_ownership'))}")
            print(f"      sheet projects : {d.get('sheet_projects')}")
            print(f"      owner         : {d.get('primary_owner')}")
            print(f"      split         : {d.get('ownership_split')}")
            print(f"      CRM leads     : {d.get('crm_lead_companies')}")
            print(f"      note          : {str(d.get('internal_note'))[:150]}")

    hdr, body = rows_of(wb["Reconciliation"])
    rec = [dict(zip(hdr, r)) for r in body]
    print()
    print("=" * 104)
    print("COMPANY ALIGNMENT ACROSS ALL 298 SHEET ROWS")
    print("=" * 104)
    print(f"  company_alignment : {dict(Counter(str(r.get('company_alignment')) for r in rec))}")
    print(f"  owning company    : "
          f"{dict(Counter(str(r.get('odoo_company')) for r in rec).most_common())}")
    print(f"  company_source    : {dict(Counter(str(r.get('company_source')) for r in rec))}")

    mis = [r for r in rec if str(r.get("company_alignment")) == "MISMATCH"]
    print()
    print(f"  MISMATCHED ROWS ({len(mis)}) - document company differs from unit owner:")
    for r in mis:
        print(f"    {r.get('project')}/{r.get('unit_no')}  client={str(r.get('client_sheet'))[:28]}")
        print(f"      unit owner      : {r.get('odoo_company')}")
        print(f"      invoice company : {r.get('odoo_invoice_companies')}")
        print(f"      payment company : {r.get('odoo_payment_companies')}")
        print(f"      note            : {str(r.get('internal_note'))[:170]}")
        print(f"      issues          : {str(r.get('issues'))[:170]}")

    # Sample of the internal_note field on normal rows
    print()
    print("=" * 104)
    print("SAMPLE internal_note ON ORDINARY ROWS")
    print("=" * 104)
    shown = 0
    for r in rec:
        if str(r.get("company_alignment")) == "OK" and r.get("internal_note"):
            print(f"  {r.get('project')}/{r.get('unit_no')}: {str(r.get('internal_note'))[:160]}")
            shown += 1
            if shown >= 4:
                break

    print()
    print("=" * 104)
    print("ODOO-ONLY UNITS BY OWNING COMPANY")
    print("=" * 104)
    hdr, body = rows_of(wb["Unreconciled"])
    # Odoo-only rows are built with a smaller dict than the reconciliation rows,
    # so map by position against this sheet's own header rather than reusing
    # the Reconciliation header.
    all_cols: list[str] = []
    for name in ("Unreconciled",):
        h, _b = rows_of(wb[name])
        for c in h:
            if c not in all_cols:
                all_cols.append(c)
    only = []
    for r in body:
        d = {all_cols[i]: v for i, v in enumerate(r) if i < len(all_cols)}
        if str(d.get("client_match")) == "NOT_IN_SHEET":
            only.append(d)
    for comp, n in Counter(str(o.get("odoo_company")) for o in only).most_common():
        blds = sorted({str(o.get("odoo_building")) for o in only
                       if str(o.get("odoo_company")) == comp})
        print(f"  {n:>5} units  {comp}")
        print(f"           buildings: {', '.join(blds)}")
    wb.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())