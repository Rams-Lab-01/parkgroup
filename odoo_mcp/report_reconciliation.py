"""Print the reconciliation workbook's defect log and clarification list."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

for _s in ("stdout", "stderr"):
    _st = getattr(sys, _s, None)
    if _st and hasattr(_st, "reconfigure"):
        try:
            _st.reconfigure(encoding="utf-8", errors="replace")
        except (OSError, ValueError):
            pass

WB = Path(__file__).parent / "reconciliation_2026-10-06.xlsx"


def rows_of(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    return rows[0], rows[1:]


def main() -> int:
    wb = load_workbook(WB, read_only=True, data_only=True)
    print("SHEETS:", wb.sheetnames)
    for ws in wb.worksheets:
        hdr, body = rows_of(ws)
        print(f"  {ws.title:<20} {len(body):>6} rows, {len(hdr)} cols")

    hdr, body = rows_of(wb["Defects_Log"])
    print()
    print("=" * 100)
    print(f"DEFECTS LOG  ({len(body)})")
    print("=" * 100)
    order = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}
    recs = [dict(zip(hdr, r)) for r in body]
    recs.sort(key=lambda d: (order.get(str(d.get("severity")), 9), str(d.get("category"))))
    for d in recs:
        print(f"  [{str(d.get('severity')):<6}] {str(d.get('category')):<34} "
              f"{str(d.get('scope')):<5} n={str(d.get('record_count')):<7} "
              f"{str(d.get('location'))}")
        print(f"           {str(d.get('description'))[:150]}")
        if d.get("example"):
            print(f"           eg: {str(d.get('example'))[:120]}")

    hdr, body = rows_of(wb["Clarifications"])
    print()
    print("=" * 100)
    print(f"CLARIFICATIONS NEEDED  ({len(body)})")
    print("=" * 100)
    for r in body:
        d = dict(zip(hdr, r))
        print(f"  {d.get('id')}  [{d.get('priority')}]  {d.get('topic')}")
        print(f"      Q      : {d.get('question')}")
        print(f"      blocks : {d.get('blocks')}")
        print(f"      assumed: {d.get('our_assumption')}")
        print()

    # Why are so many ambiguous / unreconciled?
    hdr, body = rows_of(wb["Ambiguous"])
    amb = [dict(zip(hdr, r)) for r in body]
    from collections import Counter
    print("=" * 100)
    print(f"AMBIGUOUS reasons (top 12)  [{len(amb)} rows]")
    print("=" * 100)
    for reason, n in Counter(str(a.get("reason")) for a in amb).most_common(12):
        print(f"  {n:>4}  {reason[:140]}")

    hdr, body = rows_of(wb["Unreconciled"])
    unr = [dict(zip(hdr, r)) for r in body]
    print()
    print("=" * 100)
    print(f"UNRECONCILED reasons (top 12)  [{len(unr)} rows]")
    print("=" * 100)
    for reason, n in Counter(str(u.get("reason")) for u in unr).most_common(12):
        print(f"  {n:>4}  {reason[:140]}")
    print()
    print("  by building:")
    for b, n in Counter(str(u.get("odoo_building")) for u in unr).most_common(10):
        print(f"  {n:>5}  {b}")

    # The "by building" counts above mix two different populations; split them.
    print()
    print("=" * 100)
    print("UNRECONCILED split: ODOO_ONLY vs SHEET rows")
    print("=" * 100)
    odoo_only = [u for u in unr if str(u.get("client_match")) == "NOT_IN_SHEET"]
    sheet_rows = [u for u in unr if str(u.get("client_match")) != "NOT_IN_SHEET"]
    print(f"  ODOO_ONLY (in Odoo, no sheet row) : {len(odoo_only):>5}")
    for b, n in Counter(str(u.get("odoo_building")) for u in odoo_only).most_common(10):
        print(f"      {n:>5}  {b}")
    print(f"  SHEET rows needing attention      : {len(sheet_rows):>5}")
    for b, n in Counter(str(u.get("project")) for u in sheet_rows).most_common(10):
        print(f"      {n:>5}  project {b}")

    # How many amount comparisons were actually valid?
    hdr, body = rows_of(wb["Reconciliation"])
    rec = [dict(zip(hdr, r)) for r in body]
    print()
    print("=" * 100)
    print("MONEY-COMPARISON COVERAGE (per-unit vs per-customer)")
    print("=" * 100)
    comp = Counter(str(r.get("amount_comparable")) for r in rec)
    print(f"  amount_comparable YES : {comp.get('YES', 0):>5}  (client owns exactly 1 unit)")
    print(f"  amount_comparable NO  : {comp.get('NO', 0):>5}  (multi-unit client, not comparable)")
    print(f"  collected_verdict     : {dict(Counter(str(r.get('collected_verdict')) for r in rec))}")
    print(f"  price_verdict         : {dict(Counter(str(r.get('price_verdict')) for r in rec))}")
    matched = [r for r in rec if str(r.get("collected_verdict")) == "MATCH"]
    differs = [r for r in rec if str(r.get("collected_verdict")) == "DIFFERS"]
    print(f"  collected MATCH={len(matched)}  DIFFERS={len(differs)}")
    for r in differs[:10]:
        print(f"      {r.get('project')}/{r.get('unit_no')}  client={str(r.get('client_sheet'))[:24]:<26}"
              f" sheet={r.get('collected_sheet')!s:>14}  odoo={r.get('odoo_payment_total')!s:>14}"
              f"  delta={r.get('collected_delta')}")

    # Partner matching quality
    print()
    print("=" * 100)
    print("CUSTOMER MATCH QUALITY")
    print("=" * 100)
    for v, n in Counter(str(r.get("client_match")) for r in rec).most_common():
        print(f"  {n:>5}  {v}")
    wb.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())